#!/usr/bin/env python3
"""Gives every post a real, explanatory hero image, even posts nobody made a
custom hero for (e.g. the generate-posts fallback).

For each post in posts/drafts, posts/ready (and posts/published with
--published) that has no repo image yet: take the table under "## Example"
(or the first table), put it in a spreadsheet exactly as written, render it
with LibreOffice Calc, place it on a 1600x900 canvas, and insert it right
after the "> Works in" callout. The caption says what it is: the post's own
example laid out in a spreadsheet, not a claim that anything was tested.

Run from the repo root: python3 scripts/auto-hero.py [--published]
Needs libreoffice-calc, openpyxl, pymupdf, pillow.
"""
import os
import re
import sys

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

sys.path.insert(0, os.path.dirname(__file__))
from importlib.util import spec_from_file_location, module_from_spec

_spec = spec_from_file_location("mh", os.path.join(os.path.dirname(__file__), "make-heroes.py"))
mh = module_from_spec(_spec)
_spec.loader.exec_module(mh)

HEADER_FILL = PatternFill("solid", fgColor="2E5E4E")
ALERT_FILL = PatternFill("solid", fgColor="F8D7DA")
DONE_FILL = PatternFill("solid", fgColor="E2F0D9")


def slugify(text, mx=60):
    return re.sub(r"[^a-z0-9]+", "-", text.lower()).strip("-")[:mx].rstrip("-")


def strip_md(cell):
    cell = re.sub(r"\[([^\]]+)\]\([^)]+\)", r"\1", cell)
    return re.sub(r"[*`]", "", cell).strip()


def example_table(body):
    lines = body.split("\n")
    start = next((i for i, l in enumerate(lines) if re.match(r"^##\s+Example", l)), 0)
    for i in range(start, len(lines)):
        if lines[i].startswith("|"):
            block = []
            while i < len(lines) and lines[i].startswith("|"):
                block.append(lines[i])
                i += 1
            rows = [[strip_md(c) for c in r.strip().strip("|").split("|")] for r in block]
            rows = [r for r in rows if not all(re.fullmatch(r":?-+:?", c or "-") for c in r)]
            if len(rows) >= 2:
                return rows
    return None


def build_xlsx(rows, path):
    wb = Workbook()
    ws = wb.active
    for r, row in enumerate(rows, start=1):
        for c, val in enumerate(row, start=1):
            cell = ws.cell(row=r, column=c, value=val)
            cell.number_format = "@"  # keep 02108, $18.00, 03/04/2026 exactly as written
            cell.alignment = Alignment(wrap_text=True, vertical="center")
            if r == 1:
                cell.fill, cell.font = HEADER_FILL, Font(bold=True, color="FFFFFF")
        # Shade rows the example itself calls out, the way the post's
        # conditional formatting would.
        text = " ".join(row).lower()
        if r > 1 and re.search(r"\b(overdue|reorder|duplicate|red)\b", text):
            for c in range(1, len(row) + 1):
                ws.cell(row=r, column=c).fill = ALERT_FILL
        elif r > 1 and re.search(r"\bpaid\b", text) and "unpaid" not in text:
            for c in range(1, len(row) + 1):
                ws.cell(row=r, column=c).fill = DONE_FILL
    for c in range(1, len(rows[0]) + 1):
        width = max(len(str(row[c - 1])) if c - 1 < len(row) else 0 for row in rows)
        ws.column_dimensions[ws.cell(row=1, column=c).column_letter].width = min(max(width + 2, 8), 34)
    wb.save(path)


def process(path):
    raw = open(path).read()
    m = re.match(r"^---\n([\s\S]*?)\n---\n?([\s\S]*)$", raw)
    if not m:
        return
    front, body = m.group(1), m.group(2)
    if re.search(r"^!\[[^\]]*\]\((?!https?:)", body, re.M):
        return  # already has a repo image (custom hero or screenshot)
    rows = example_table(body)
    if not rows:
        print(f"skip {path}: no example table")
        return
    file = os.path.basename(path)[:-3]
    title = re.search(r"^title:\s*(.+)$", front, re.M).group(1).strip()
    xlsx = os.path.join(mh.rp.TMP, f"{file}.xlsx")
    build_xlsx(rows, xlsx)
    img = mh.render(xlsx, "Sheet", name=file)
    out = f"images/{file}/{slugify(title, 50)}-example.png"
    mh.save(out, img)
    cols = ", ".join(rows[0][:4])
    alt = f"Spreadsheet example with columns {cols}"
    alt = " ".join(alt.split()[:16])
    line = f"![{alt}]({out}) *The example from this post laid out in a spreadsheet. Rendered with LibreOffice Calc.*"
    lines = body.split("\n")
    i = next((k for k, l in enumerate(lines) if l.startswith("> ")), None)
    if i is None:
        i = next(k for k, l in enumerate(lines) if l.strip())  # after the first paragraph
        while i < len(lines) and lines[i].strip():
            i += 1
    else:
        while i < len(lines) and lines[i].startswith(">"):
            i += 1
    lines[i:i] = ["", line]
    open(path, "w").write(f"---\n{front}\n---\n" + "\n".join(lines))
    print(f"hero {out} -> {path}")


def main():
    dirs = ["posts/drafts", "posts/ready"] + (["posts/published"] if "--published" in sys.argv else [])
    for d in dirs:
        if os.path.isdir(d):
            for f in sorted(os.listdir(d)):
                if f.endswith(".md"):
                    process(os.path.join(d, f))


if __name__ == "__main__":
    main()
