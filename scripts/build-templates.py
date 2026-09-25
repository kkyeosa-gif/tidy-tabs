#!/usr/bin/env python3
"""Builds the downloadable .xlsx/.csv templates in templates/.

Run from the repo root: python3 scripts/build-templates.py
Needs openpyxl (pip install openpyxl). Every template is plain formulas
(IF, SUMIFS, TODAY, TEXT) so it opens the same in Excel and after
File > Import in Google Sheets. Re-run after editing; the files are
committed so readers download exactly what was checked.
"""
import csv
from datetime import date
from pathlib import Path

from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

OUT = Path("templates")
USD = '"$"#,##0.00'
US_DATE = "mm/dd/yyyy"
HEADER_FILL = PatternFill("solid", fgColor="2E5E4E")
HEADER_FONT = Font(bold=True, color="FFFFFF")
ALERT_FILL = PatternFill("solid", fgColor="F8D7DA")
DONE_FILL = PatternFill("solid", fgColor="E2F0D9")
FORMULA_FILL = PatternFill("solid", fgColor="EDEDED")


def header(ws, names, widths):
    ws.append(names)
    for i, (cell, w) in enumerate(zip(ws[1], widths), start=1):
        cell.fill, cell.font = HEADER_FILL, HEADER_FONT
        cell.alignment = Alignment(vertical="center", wrap_text=True)
        ws.column_dimensions[cell.column_letter].width = w
    ws.row_dimensions[1].height = 30
    ws.freeze_panes = "A2"


def notes_sheet(wb, lines):
    ws = wb.create_sheet("How to use")
    ws.column_dimensions["A"].width = 100
    for line in lines:
        ws.append([line])
    ws["A1"].font = Font(bold=True, size=13)


def inventory():
    wb = Workbook()
    items = wb.active
    items.title = "Items"
    header(items, ["SKU", "Item", "Category", "Unit cost", "Sale price", "Starting qty",
                   "Qty in", "Qty out", "On hand", "Reorder at", "Status", "Stock value (cost)"],
           [10, 30, 14, 11, 11, 11, 9, 9, 9, 11, 13, 16])
    rows = [
        ("CND-LAV8", "Lavender soy candle, 8 oz", "Candles", 4.10, 18.00, 40, 10),
        ("CND-CED8", "Cedar soy candle, 8 oz", "Candles", 4.25, 18.00, 25, 10),
        ("SOP-OAT", "Oatmeal goat milk soap bar", "Soap", 1.85, 8.00, 60, 20),
        ("SOP-CHR", "Charcoal soap bar", "Soap", 2.05, 8.50, 30, 20),
        ("EAR-HOOP", "Brass hoop earrings", "Jewelry", 3.40, 24.00, 15, 5),
        ("CRD-BDAY", "Letterpress birthday card", "Cards", 0.95, 6.00, 80, 25),
        ("PKG-BOX6", "Kraft mailer box 6x6x4", "Supplies", 0.62, 0, 100, 40),
    ]
    for r, (sku, name, cat, cost, price, start, reorder) in enumerate(rows, start=2):
        items.append([sku, name, cat, cost, price, start,
                      f"=SUMIFS('Stock log'!$C:$C,'Stock log'!$B:$B,A{r})",
                      f"=SUMIFS('Stock log'!$D:$D,'Stock log'!$B:$B,A{r})",
                      f"=F{r}+G{r}-H{r}", reorder,
                      f'=IF(A{r}="","",IF(I{r}<=J{r},"REORDER","OK"))',
                      f"=I{r}*D{r}"])
        for col in "DEL":
            items[f"{col}{r}"].number_format = USD
        for col in "GHIKL":
            items[f"{col}{r}"].fill = FORMULA_FILL
    # Ranges run to row 500 (not just the sample rows) so a product added by
    # copying the last row down is counted, offered in the dropdown, and
    # highlighted without editing any formula. The total sits beside the
    # table, not under it, for the same reason.
    last = 500
    items.column_dimensions["N"].width = 18
    items.column_dimensions["O"].width = 14
    items["N1"], items["O1"] = "Total stock value", f"=SUM(L2:L{last})"
    items["N1"].font = Font(bold=True)
    items["O1"].number_format = USD
    items.conditional_formatting.add(f"K2:K{last}", FormulaRule(formula=[f'$K2="REORDER"'], fill=ALERT_FILL))

    log = wb.create_sheet("Stock log")
    header(log, ["Date", "SKU", "Qty in", "Qty out", "Note"], [12, 12, 9, 9, 40])
    entries = [
        (date(2026, 9, 1), "CND-LAV8", 24, 0, "Poured batch #14"),
        (date(2026, 9, 3), "CND-LAV8", 0, 12, "Etsy orders 9/1-9/3"),
        (date(2026, 9, 5), "SOP-OAT", 0, 18, "Saturday farmers market"),
        (date(2026, 9, 5), "CND-CED8", 0, 9, "Saturday farmers market"),
        (date(2026, 9, 8), "SOP-CHR", 0, 14, "Wholesale order, Main St. Gift Shop"),
        (date(2026, 9, 10), "EAR-HOOP", 0, 11, "Etsy orders"),
        (date(2026, 9, 12), "PKG-BOX6", 0, 64, "Shipped Etsy orders"),
        (date(2026, 9, 15), "CRD-BDAY", 0, 30, "Craft fair"),
    ]
    for e in entries:
        log.append(list(e))
        log.cell(row=log.max_row, column=1).number_format = US_DATE
    sku_list = DataValidation(type="list", formula1=f"=Items!$A$2:$A${last}", allow_blank=True)
    log.add_data_validation(sku_list)
    sku_list.add("B2:B1000")

    notes_sheet(wb, [
        "Tidy Tabs: Small Craft Business Inventory Tracker",
        "",
        "1. Items tab: one row per product or supply. Type your own SKU, cost, price, starting count, and reorder point.",
        "2. Stock log tab: add a row every time stock comes in (Qty in) or goes out (Qty out). Pick the SKU from the dropdown.",
        "3. On hand, Status, and Stock value update by themselves. Status turns red when On hand is at or below Reorder at.",
        "4. Don't type over the gray formula columns (Qty in, Qty out, On hand, Status, Stock value).",
        "5. To add a product, copy the last Items row down so the formulas come with it. Total stock value is in cell O1.",
        "",
        "Sample data is fictional. Delete the Stock log rows and change the Items rows to start fresh.",
        "Google Sheets: File > Import > Upload this .xlsx > Insert new sheet(s) or Replace spreadsheet.",
    ])
    wb.save(OUT / "tidy-tabs-craft-inventory-tracker.xlsx")


def projects():
    wb = Workbook()
    ws = wb.active
    ws.title = "Projects"
    header(ws, ["Client", "Project", "Start date", "Due date", "Status", "Billing", "Hours", "Rate",
                "Flat fee", "Amount", "Invoice sent", "Paid on", "Days left", "Payment"],
           [18, 28, 12, 12, 13, 10, 8, 9, 10, 12, 12, 12, 10, 12])
    rows = [
        ("Harbor Dental", "Spring newsletter design", date(2026, 9, 2), date(2026, 9, 19), "Done", "Flat", None, None, 650, date(2026, 9, 19), date(2026, 9, 24)),
        ("Pine & Co. Realty", "Listing photo edits (batch 3)", date(2026, 9, 8), date(2026, 9, 26), "In progress", "Hourly", 6.5, 55, None, None, None),
        ("Lakeside Yoga", "Website copy refresh", date(2026, 9, 15), date(2026, 10, 10), "In progress", "Hourly", 3, 60, None, None, None),
        ("Main St. Gift Shop", "Holiday flyer", date(2026, 9, 22), date(2026, 10, 3), "Not started", "Flat", None, None, 300, None, None),
        ("Harbor Dental", "Social media templates", date(2026, 8, 18), date(2026, 9, 5), "Done", "Flat", None, None, 480, date(2026, 9, 5), None),
    ]
    for r, (client, proj, start, due, status, billing, hours, rate, flat, sent, paid) in enumerate(rows, start=2):
        ws.append([client, proj, start, due, status, billing, hours, rate, flat,
                   f'=IF(F{r}="Hourly",G{r}*H{r},I{r})', sent, paid,
                   f'=IF(OR(D{r}="",E{r}="Done"),"",D{r}-TODAY())',
                   f'=IF(L{r}<>"","Paid",IF(K{r}="","Not invoiced",IF(TODAY()-K{r}>30,"Overdue","Waiting")))'])
        for col in "CDKL":
            ws[f"{col}{r}"].number_format = US_DATE
        for col in "HIJ":
            ws[f"{col}{r}"].number_format = USD
    last = 200
    status = DataValidation(type="list", formula1='"Not started,In progress,Waiting on client,Done"', allow_blank=True)
    billing = DataValidation(type="list", formula1='"Hourly,Flat"', allow_blank=True)
    ws.add_data_validation(status)
    ws.add_data_validation(billing)
    status.add(f"E2:E{last}")
    billing.add(f"F2:F{last}")
    ws.conditional_formatting.add(f"A2:N{last}", FormulaRule(formula=['$E2="Done"'], fill=DONE_FILL))
    ws.conditional_formatting.add(f"M2:M{last}", FormulaRule(formula=['AND(ISNUMBER($M2),$M2<0)'], fill=ALERT_FILL))
    ws.conditional_formatting.add(f"N2:N{last}", FormulaRule(formula=['$N2="Overdue"'], fill=ALERT_FILL))

    summary = wb.create_sheet("Summary")
    summary.column_dimensions["A"].width = 28
    summary.column_dimensions["B"].width = 14
    for label, formula in [
        ("Open projects", '=COUNTIFS(Projects!A2:A200,"<>",Projects!E2:E200,"<>Done")'),
        ("Billed, not paid ($)", '=SUMIFS(Projects!J2:J200,Projects!K2:K200,"<>",Projects!L2:L200,"")'),
        ("Paid ($)", '=SUMIFS(Projects!J2:J200,Projects!L2:L200,"<>")'),
        ("Invoices overdue (30+ days)", '=COUNTIF(Projects!N2:N200,"Overdue")'),
    ]:
        summary.append([label, formula])
    for r in (2, 3):
        summary[f"B{r}"].number_format = USD

    notes_sheet(wb, [
        "Tidy Tabs: Freelance Project Tracker",
        "",
        "1. One row per project. Pick Status and Billing from the dropdowns.",
        "2. Hourly jobs: fill Hours and Rate. Flat-fee jobs: fill Flat fee. Amount picks the right one.",
        "3. Type the date you sent the invoice in Invoice sent, and the date the money arrived in Paid on.",
        "4. Days left counts down to the due date (negative = late) and hides itself once Status is Done.",
        "5. Payment shows Not invoiced / Waiting / Overdue (more than 30 days after you invoiced) / Paid.",
        "6. Days left and Payment use TODAY(), so they change every day you open the file.",
        "",
        "Sample clients are fictional. Delete rows 2-6 on the Projects tab to start fresh.",
        "Google Sheets: File > Import > Upload this .xlsx.",
    ])
    wb.save(OUT / "tidy-tabs-freelance-project-tracker.xlsx")


ZIPS = [
    ("Boston", "MA", "02108", "02108-1522"),
    ("Hoboken", "NJ", "07030", "07030-3707"),
    ("Holtsville", "NY", "00501", "00501-0001"),
    ("Burlington", "VT", "05401", "05401-4417"),
    ("Chicago", "IL", "60601", "60601-3014"),
    ("Las Vegas", "NV", "89101", "89101-5019"),
]


def zip_codes():
    wb = Workbook()
    ws = wb.active
    ws.title = "Fixed (text)"
    header(ws, ["City", "State", "ZIP (text)", "ZIP+4 (text)"], [14, 7, 12, 14])
    for city, st, z, z4 in ZIPS:
        ws.append([city, st, z, z4])
    for row in ws.iter_rows(min_row=2, max_row=1000, min_col=3, max_col=4):
        for c in row:
            c.number_format = "@"

    broken = wb.create_sheet("Broken (numbers)")
    header(broken, ["City", "State", "ZIP as typed", "Repair with TEXT", "Repair ZIP+4 (9 digits)"], [14, 7, 14, 18, 22])
    for r, (city, st, z, z4) in enumerate(ZIPS, start=2):
        broken.append([city, st, int(z), f'=TEXT(C{r},"00000")', f'=TEXT({int(z4.replace("-", ""))},"00000-0000")'])

    notes_sheet(wb, [
        "Tidy Tabs: US ZIP Code Leading Zeros sample",
        "",
        "Fixed (text): the ZIP columns are formatted as Text (@) BEFORE the codes were entered, so 02108 stays 02108.",
        "Broken (numbers): the same ZIPs stored as numbers, which is what happens when you type or paste them into a General cell.",
        'Column D rebuilds a 5-digit ZIP with =TEXT(C2,"00000"). Copy it, then Paste Special > Values to keep the result.',
        "Column E shows the ZIP+4 version with the format 00000-0000.",
        "",
        "The 5-digit ZIPs are real; the +4 extensions are made up for practice.",
        "The same six ZIP codes are in tidy-tabs-zip-codes-sample.csv for practicing a CSV import.",
    ])
    wb.save(OUT / "tidy-tabs-zip-codes-leading-zeros.xlsx")

    with open(OUT / "tidy-tabs-zip-codes-sample.csv", "w", newline="") as f:
        w = csv.writer(f)
        w.writerow(["City", "State", "ZIP", "ZIP+4"])
        for row in ZIPS:
            w.writerow(row)


def print_one_page():
    wb = Workbook()
    ws = wb.active
    ws.title = "Price list"
    cols = ["SKU", "Item", "Category", "Size", "Color", "Unit cost", "Wholesale", "Retail",
            "Margin %", "On hand", "Reorder at", "Supplier", "Lead time (days)", "Last ordered"]
    header(ws, cols, [10, 30, 12, 10, 10, 10, 11, 10, 10, 9, 12, 14, 10, 13])
    colors = ["Natural", "Sage", "Navy", "Rust", "Cream"]
    sizes = ["4 oz", "8 oz", "12 oz"]
    for i in range(1, 49):
        r = i + 1
        cost = round(1.5 + (i % 7) * 0.65, 2)
        ws.append([f"ITM-{i:03d}", f"Sample product {i}", ["Candles", "Soap", "Cards", "Jewelry"][i % 4],
                   sizes[i % 3], colors[i % 5], cost, f"=ROUND(F{r}*2,2)", f"=ROUND(F{r}*{3 + (i % 3) * 0.5},2)",
                   f"=(H{r}-F{r})/H{r}", 10 + (i * 7) % 50, 15, f"Supplier {chr(65 + i % 6)}",
                   7 + (i % 4) * 7, date(2026, 8, 1 + i % 28)])
        for col in "FGH":
            ws[f"{col}{r}"].number_format = USD
        ws[f"I{r}"].number_format = "0%"
        ws[f"N{r}"].number_format = US_DATE
    # The settings the post walks through, already applied: US Letter,
    # landscape, fit all columns on one page wide, header row repeats.
    ws.page_setup.paperSize = ws.PAPERSIZE_LETTER
    ws.page_setup.orientation = "landscape"
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 1
    ws.print_title_rows = "1:1"
    ws.page_margins.left = ws.page_margins.right = 0.25
    ws.page_margins.top = ws.page_margins.bottom = 0.5
    ws.print_options.horizontalCentered = True
    wb.save(OUT / "tidy-tabs-print-one-page-us-letter.xlsx")


ADDRESSES = [
    ("Boston, MA 02108", "Harbor Dental"),
    ("Hoboken, NJ 07030", "Pine & Co. Realty"),
    ("Holtsville, NY 00501", "Main St. Gift Shop"),
    ("Burlington, VT 05401", "Lakeside Yoga"),
    ("Chicago, IL 60601", "Northside Print Co."),
    ("Las Vegas, NV 89101", "Desert Bloom Candles"),
    ("Salt Lake City, UT 84101", "Wasatch Bike Repair"),
    ("St. Louis, MO 63101", "Gateway Bakery"),
]


def split_city_state_zip():
    wb = Workbook()
    ws = wb.active
    ws.title = "Split"
    header(ws, ["Customer", "City, ST ZIP", "City", "State", "ZIP"], [22, 28, 18, 8, 9])
    for r, (addr, name) in enumerate(ADDRESSES, start=2):
        ws.append([name, addr,
                   f'=LEFT(B{r},FIND(",",B{r})-1)',
                   f'=MID(B{r},FIND(",",B{r})+2,2)',
                   f'=RIGHT(B{r},5)'])
        for col in "CDE":
            ws[f"{col}{r}"].fill = FORMULA_FILL
    notes_sheet(wb, [
        "Tidy Tabs: Split City, ST ZIP into three columns",
        "",
        'Type or paste addresses like "Boston, MA 02108" in column B.',
        "City, State, and ZIP (gray) fill in by formula. Copy a gray row down for more addresses.",
        'City = LEFT(B2,FIND(",",B2)-1): everything before the comma.',
        'State = MID(B2,FIND(",",B2)+2,2): the 2 letters after the comma and space.',
        "ZIP = RIGHT(B2,5): the last 5 characters, returned as text, so 02108 keeps its 0.",
        "Works for 5-digit ZIPs. For ZIP+4 (02108-1522), use RIGHT(B2,10).",
        "To keep only the results, copy C:E and Paste Special > Values.",
        "",
        "Customer names are fictional; the city/state/ZIP combinations are real.",
    ])
    wb.save(OUT / "tidy-tabs-split-city-state-zip.xlsx")


if __name__ == "__main__":
    OUT.mkdir(exist_ok=True)
    inventory()
    projects()
    zip_codes()
    print_one_page()
    split_city_state_zip()
    print("\n".join(sorted(p.name for p in OUT.iterdir())))
