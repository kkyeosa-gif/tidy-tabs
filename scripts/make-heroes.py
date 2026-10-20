#!/usr/bin/env python3
"""Builds each post's hero image: a real LibreOffice render of the post's
template (or a before/after of the real print output), placed on a 1600x900
(16:9) canvas. The hero is the first image in the post, so Blogger uses it
as og:image, and Google Discover wants >=1200px-wide 16:9 images that are
"relevant and representative" (see tasks/briefs.md 2026-09-26).

Run from the repo root: python3 scripts/make-heroes.py
Needs libreoffice-calc, openpyxl, pymupdf, pillow.
"""
import io
import os
import sys

import pymupdf
from openpyxl import load_workbook
from PIL import Image, ImageFilter

from importlib.util import spec_from_file_location, module_from_spec

_spec = spec_from_file_location("rp", os.path.join(os.path.dirname(__file__), "render-previews.py"))
rp = module_from_spec(_spec)
_spec.loader.exec_module(rp)

W, H, PAD = 1600, 900, 50
BG = (244, 246, 243)


def render(src, sheet, hide_cols="", name="hero", zoom=3.0, widths=None):
    """PDF-render one sheet (optionally hiding columns) and return a cropped PIL image."""
    wb = load_workbook(src)
    wb.move_sheet(sheet, -wb.index(wb[sheet]))
    for ws in wb.worksheets:
        ws.page_setup.paperSize = ws.PAPERSIZE_LETTER
        ws.page_setup.orientation = "landscape"
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.page_setup.fitToWidth = ws.page_setup.fitToHeight = 1
        ws.print_options.gridLines = True
    for col in hide_cols:
        wb[sheet].column_dimensions[col].hidden = True
    for col, width in (widths or {}).items():
        wb[sheet].column_dimensions[col].width = width
    path = f"{rp.TMP}/{name}.xlsx"
    wb.save(path)
    page = pymupdf.open(rp.to_pdf(path))[0]
    rects = [b[:4] for b in page.get_text("blocks")] + [d["rect"] for d in page.get_drawings()]
    clip = pymupdf.Rect(rects[0])
    for r in rects[1:]:
        clip |= pymupdf.Rect(r)
    pix = page.get_pixmap(matrix=pymupdf.Matrix(zoom, zoom), clip=clip + (-6, -6, 6, 6))
    return Image.open(io.BytesIO(pix.tobytes("png"))).convert("RGB")


def card(img, max_w, max_h):
    scale = min(max_w / img.width, max_h / img.height)
    img = img.resize((int(img.width * scale), int(img.height * scale)), Image.LANCZOS)
    shadow = Image.new("RGBA", (img.width + 24, img.height + 24), (0, 0, 0, 0))
    shadow.paste((0, 0, 0, 60), (12, 16, img.width + 12, img.height + 16))
    shadow = shadow.filter(ImageFilter.GaussianBlur(8))
    shadow.paste(img, (12, 12))
    return shadow


def save(out, *images):
    canvas = Image.new("RGB", (W, H), BG)
    n = len(images)
    slot_w = (W - PAD * (n + 1)) // n
    for i, img in enumerate(images):
        c = card(img, slot_w, H - 2 * PAD)
        x = PAD + i * (slot_w + PAD) + (slot_w - c.width) // 2
        canvas.paste(c, (x, (H - c.height) // 2), c)
    os.makedirs(os.path.dirname(out), exist_ok=True)
    canvas.save(out, optimize=True)
    print("wrote", out, os.path.getsize(out) // 1024, "KB")


def heroes_2026_10_06():
    t = "templates/"
    d = "images/2026-10-06-team-0%d-%s/"
    save(d % (1, "loan-payment-calculator-pmt-excel-google-sheets") + "loan-payment-calculator-template.png",
         render(t + "tidy-tabs-loan-payment-calculator.xlsx", "Loans", name="loans"))
    save(d % (2, "percent-change-month-over-month-excel-google-sheets") + "percent-change-template.png",
         render(t + "tidy-tabs-percent-change.xlsx", "Sales", name="pctchange"))
    save(d % (3, "rank-top-customers-excel-google-sheets") + "rank-top-customers-template.png",
         render(t + "tidy-tabs-rank-top-customers.xlsx", "Customers", name="rank"))
    save(d % (4, "invoice-number-generator-text-excel-google-sheets") + "invoice-number-generator-template.png",
         render(t + "tidy-tabs-invoice-number-generator.xlsx", "Invoices", name="invnum"))
    save(d % (5, "combine-address-columns-textjoin-excel-google-sheets") + "combine-address-columns-textjoin-template.png",
         render(t + "tidy-tabs-combine-address-columns.xlsx", "Addresses", hide_cols="H", name="addr"))


def heroes_2026_10_07():
    t = "templates/"
    d = "images/2026-10-07-team-0%d-%s/"
    save(d % (1, "sale-price-discount-calculator-excel-google-sheets") + "sale-price-discount-calculator-template.png",
         render(t + "tidy-tabs-sale-price-discount-calculator.xlsx", "Sale prices", hide_cols="F", name="saleprice"))
    save(d % (2, "shipping-cost-weight-tier-lookup-excel-google-sheets") + "shipping-weight-tier-lookup-template.png",
         render(t + "tidy-tabs-shipping-weight-tier-lookup.xlsx", "Orders", hide_cols="DE", name="shiptier"))
    save(d % (3, "reorder-point-calculator-excel-google-sheets") + "reorder-point-calculator-template.png",
         render(t + "tidy-tabs-reorder-point-calculator.xlsx", "Reorder", hide_cols="I", name="reorder"))
    save(d % (4, "freelance-hourly-rate-calculator-excel-google-sheets") + "freelance-hourly-rate-calculator-template.png",
         render(t + "tidy-tabs-freelance-hourly-rate-calculator.xlsx", "Hourly rate", name="hourly"))
    save(d % (5, "convert-text-to-numbers-dollar-signs-excel-google-sheets") + "convert-text-to-numbers-template.png",
         render(t + "tidy-tabs-convert-text-to-numbers.xlsx", "Pasted amounts", hide_cols="EF", name="text2num"))


def heroes_2026_10_08():
    t = "templates/"
    d = "images/2026-10-08-team-0%d-%s/"
    save(d % (1, "sales-tax-calculator-excel-google-sheets") + "sales-tax-calculator-template.png",
         render(t + "tidy-tabs-sales-tax-calculator.xlsx", "Sales tax", name="salestax"))
    save(d % (2, "sum-expenses-by-category-sumif-excel-google-sheets") + "sum-expenses-by-category-template.png",
         render(t + "tidy-tabs-sum-expenses-by-category.xlsx", "Expenses", name="sumexp"))
    save(d % (3, "round-prices-nearest-nickel-99-excel-google-sheets") + "round-prices-nearest-nickel-99-template.png",
         render(t + "tidy-tabs-round-prices-nearest-nickel-99.xlsx", "Prices", name="roundprices"))
    save(d % (4, "remove-non-breaking-spaces-clean-substitute-excel-google-sheets") + "remove-non-breaking-spaces-template.png",
         render(t + "tidy-tabs-remove-non-breaking-spaces.xlsx", "Pasted names", hide_cols="DIJKLM", name="nbsp"))
    save(d % (5, "year-to-date-sales-total-sumifs-excel-google-sheets") + "year-to-date-sales-total-template.png",
         render(t + "tidy-tabs-year-to-date-sales-total.xlsx", "Sales", name="ytd"))


def heroes_2026_10_09():
    t = "templates/"
    d = "images/2026-10-09-team-0%d-%s/"
    save(d % (1, "quantity-price-tiers-lookup-excel-google-sheets") + "quantity-price-tiers-lookup-template.png",
         render(t + "tidy-tabs-quantity-price-tiers.xlsx", "Orders", name="tiers"))
    save(d % (2, "weighted-average-cost-sumproduct-excel-google-sheets") + "weighted-average-cost-summary-template.png",
         render(t + "tidy-tabs-weighted-average-cost.xlsx", "Summary", name="wavg"))
    save(d % (3, "last-order-date-per-customer-maxifs-excel-google-sheets") + "last-order-date-per-customer-template.png",
         render(t + "tidy-tabs-last-order-date-maxifs.xlsx", "Summary", name="maxifs"))
    save(d % (4, "count-orders-by-month-countifs-excel-google-sheets") + "count-orders-by-month-template.png",
         render(t + "tidy-tabs-count-orders-by-month.xlsx", "Monthly", name="bymonth"))
    save(d % (5, "average-order-value-by-channel-averageifs-excel-google-sheets") + "average-order-value-by-channel-template.png",
         render(t + "tidy-tabs-average-order-value-by-channel.xlsx", "By channel", name="aov"))


def etsy_only_copy(src):
    """Temp copy of the subtotal template with every non-Etsy row hidden (what the filter does)."""
    wb = load_workbook(src)
    ws = wb["Sales"]
    for r in range(2, ws.max_row + 1):
        if ws.cell(r, 2).value != "Etsy":
            ws.row_dimensions[r].hidden = True
    path = f"{rp.TMP}/subtotal-etsy.xlsx"
    wb.save(path)
    return path


def heroes_2026_10_10():
    t = "templates/"
    d = "images/2026-10-10-team-0%d-%s/"
    etsy = etsy_only_copy(t + "tidy-tabs-subtotal-filtered-sales.xlsx")
    save(d % (1, "sum-visible-rows-only-subtotal-excel-google-sheets") + "subtotal-filtered-sales-template.png",
         render(etsy, "Sales", name="subsales"),
         render(etsy, "Totals", hide_cols="B", name="subtotals"))
    save(d % (2, "rolling-3-month-average-sales-excel-google-sheets") + "rolling-average-sales-template.png",
         render(t + "tidy-tabs-rolling-average-sales.xlsx", "Monthly sales", hide_cols="DE", name="rolling"))
    save(d % (3, "prorate-first-month-subscription-excel-google-sheets") + "prorate-partial-month-template.png",
         render(t + "tidy-tabs-prorate-partial-month.xlsx", "Proration", hide_cols="F", name="prorate"))
    save(d % (4, "two-way-rate-card-lookup-index-match-excel-google-sheets") + "two-way-rate-card-template.png",
         render(t + "tidy-tabs-two-way-rate-card.xlsx", "Quote", name="quote"))
    save(d % (5, "in-cell-bar-chart-rept-excel-google-sheets") + "in-cell-bars-template.png",
         render(t + "tidy-tabs-in-cell-bars.xlsx", "Units sold", hide_cols="DE", name="bars"))


def heroes_2026_10_11():
    t = "templates/"
    d = "images/2026-10-11-team-0%d-%s/"
    save(d % (1, "quarterly-sales-totals-sumifs-excel-google-sheets") + "quarterly-sales-totals-by-quarter-template.png",
         render(t + "tidy-tabs-quarterly-sales-totals.xlsx", "By quarter", name="quarters"))
    save(d % (2, "find-missing-invoice-numbers-countif-excel-google-sheets") + "missing-invoice-numbers-check-template.png",
         render(t + "tidy-tabs-missing-invoice-numbers.xlsx", "Check", name="missinginv", widths={"E": 46}))
    save(d % (3, "sum-if-cell-contains-text-wildcard-sumif-excel-google-sheets") + "sumif-contains-text-totals-template.png",
         render(t + "tidy-tabs-sumif-contains-text.xlsx", "Totals", name="sumifwild"))
    save(d % (4, "split-invoice-into-installments-excel-google-sheets") + "invoice-installments-payment-plan-template.png",
         render(t + "tidy-tabs-invoice-installments.xlsx", "Payment plan", name="installments", widths={"A": 34}))
    save(d % (5, "marketplace-fee-net-payout-calculator-excel-google-sheets") + "marketplace-net-payout-calculator-template.png",
         render(t + "tidy-tabs-marketplace-net-payout.xlsx", "Net payout", name="netpayout"))


def heroes_2026_10_12():
    t = "templates/"
    d = "images/2026-10-12-team-0%d-%s/"
    save(d % (1, "compare-two-lists-find-missing-countif-excel-google-sheets") + "compare-two-lists-find-missing-template.png",
         render(t + "tidy-tabs-compare-two-lists.xlsx", "Lists", name="lists"))
    save(d % (2, "tiered-sales-commission-sumproduct-excel-google-sheets") + "tiered-sales-commission-template.png",
         render(t + "tidy-tabs-tiered-commission.xlsx", "Commission", name="commission"))
    save(d % (3, "monthly-sales-goal-progress-excel-google-sheets") + "monthly-sales-goal-progress-template.png",
         render(t + "tidy-tabs-sales-goal-progress.xlsx", "Goal", name="goal"))
    save(d % (4, "batch-cost-per-unit-excel-google-sheets") + "batch-cost-per-unit-template.png",
         render(t + "tidy-tabs-batch-cost-per-unit.xlsx", "Batch cost", name="batch"))
    save(d % (5, "percent-of-total-product-sales-excel-google-sheets") + "percent-of-total-sales-by-product-template.png",
         render(t + "tidy-tabs-percent-of-total-sales.xlsx", "Sales by product", name="pct"))


def heroes_2026_10_13():
    t = "templates/"
    d = "images/2026-10-13-team-0%d-%s/"
    save(d % (1, "invoice-due-date-payment-terms-excel-google-sheets") + "invoice-due-date-payment-terms-template.png",
         render(t + "tidy-tabs-invoice-due-dates.xlsx", "Due dates", name="duedates"))
    save(d % (2, "sales-by-day-of-week-sumif-excel-google-sheets") + "sales-by-day-of-week-template.png",
         render(t + "tidy-tabs-sales-by-weekday.xlsx", "By weekday", name="weekday"))
    save(d % (3, "repeat-customer-rate-countif-excel-google-sheets") + "repeat-customer-rate-template.png",
         render(t + "tidy-tabs-repeat-customers.xlsx", "Customers", name="repeat"))
    save(d % (4, "cheapest-supplier-cost-per-ounce-index-match-excel-google-sheets") + "cheapest-supplier-cost-per-ounce-template.png",
         render(t + "tidy-tabs-supplier-price-compare.xlsx", "Compare", name="supplier"))
    save(d % (5, "split-sku-code-into-parts-left-mid-find-excel-google-sheets") + "split-sku-code-into-parts-template.png",
         render(t + "tidy-tabs-split-sku-code.xlsx", "SKUs", name="sku"))


def heroes_2026_10_14():
    t = "templates/"
    d = "images/2026-10-14-team-0%d-%s/"
    save(d % (3, "count-listing-characters-tags-len-substitute-excel-google-sheets") + "listing-length-check-template.png",
         render(t + "tidy-tabs-listing-length-check.xlsx", "Listings", name="listinglen"))
    save(d % (1, "convert-ounces-to-pounds-and-ounces-int-mod-excel-google-sheets") + "ounces-to-pounds-and-ounces-template.png",
         render(t + "tidy-tabs-ounces-to-pounds-ounces.xlsx", "Weights", name="ounces"))
    save(d % (2, "convert-time-zones-client-calls-excel-google-sheets") + "time-zone-converter-client-calls-template.png",
         render(t + "tidy-tabs-time-zone-converter.xlsx", "Calls", name="timezone"))
    save(d % (4, "last-number-in-column-lookup-excel-google-sheets") + "last-number-in-column-lookup-template.png",
         render(t + "tidy-tabs-last-number-in-column.xlsx", "Log", name="lastlog"),
         render(t + "tidy-tabs-last-number-in-column.xlsx", "Summary", name="lastsum"))
    save(d % (5, "count-market-saturdays-in-month-networkdays-intl-excel-google-sheets") + "count-market-saturdays-by-month-template.png",
         render(t + "tidy-tabs-count-weekdays-in-month.xlsx", "Months", name="saturdays"))


def heroes_2026_10_15():
    t = "templates/"
    d = "images/2026-10-15-team-0%d-%s/"
    save(d % (1, "average-days-to-get-paid-excel-google-sheets") + "days-to-get-paid-template.png",
         render(t + "tidy-tabs-days-to-get-paid.xlsx", "Invoices", name="daystopaid"))
    save(d % (2, "early-payment-discount-2-10-net-30-excel-google-sheets") + "early-payment-discount-template.png",
         render(t + "tidy-tabs-early-payment-discount.xlsx", "Discounts", name="earlydisc"))
    save(d % (3, "sell-through-rate-inventory-excel-google-sheets") + "sell-through-rate-template.png",
         render(t + "tidy-tabs-sell-through-rate.xlsx", "Products", name="sellthru"))
    save(d % (4, "how-many-boxes-needed-roundup-excel-google-sheets") + "boxes-needed-template.png",
         render(t + "tidy-tabs-boxes-needed.xlsx", "Orders", name="boxes"))
    save(d % (5, "median-order-value-vs-average-excel-google-sheets") + "median-order-value-template.png",
         render(t + "tidy-tabs-median-order-value.xlsx", "Orders", name="median"))


def short_log_copy(src, last_row=20):
    """Temp copy of the week-ending Log tab with the empty rows below the data hidden (keeps the hero 16:9 readable)."""
    wb = load_workbook(src)
    ws = wb["Log"]
    for r in range(last_row + 1, ws.max_row + 1):
        ws.row_dimensions[r].hidden = True
    path = f"{rp.TMP}/weekending-short.xlsx"
    wb.save(path)
    return path


def heroes_2026_10_16():
    t = "templates/"
    d = "images/2026-10-16-team-0%d-%s/"
    save(d % (1, "week-ending-friday-date-weekday-excel-google-sheets") + "week-ending-date-template.png",
         render(short_log_copy(t + "tidy-tabs-week-ending-date.xlsx"), "Log", name="weekending"))
    save(d % (2, "free-shipping-threshold-amount-to-add-if-and-excel-google-sheets") + "free-shipping-threshold-template.png",
         render(t + "tidy-tabs-free-shipping-threshold.xlsx", "Orders", name="freeship"))
    save(d % (3, "forecast-next-month-sales-forecast-linear-excel-google-sheets") + "forecast-next-month-sales-template.png",
         render(t + "tidy-tabs-forecast-next-month.xlsx", "Sales", name="forecast"))
    save(d % (4, "client-tenure-years-months-datedif-excel-google-sheets") + "client-tenure-years-months-template.png",
         render(t + "tidy-tabs-client-tenure.xlsx", "Clients", name="tenure"))
    save(d % (5, "find-incomplete-customer-rows-countblank-excel-google-sheets") + "incomplete-customer-rows-template.png",
         render(t + "tidy-tabs-incomplete-rows.xlsx", "Customers", name="incomplete"))


def short_sheet_copy(src, sheet, last_row, name):
    """Temp copy of a template tab with the empty rows below the data hidden (keeps the hero 16:9 readable)."""
    wb = load_workbook(src)
    ws = wb[sheet]
    for r in range(last_row + 1, ws.max_row + 1):
        ws.row_dimensions[r].hidden = True
    path = f"{rp.TMP}/{name}-short.xlsx"
    wb.save(path)
    return path


def heroes_2026_10_17():
    t = "templates/"
    d = "images/2026-10-17-team-0%d-%s/"
    save(d % (1, "find-overlapping-bookings-date-ranges-countifs-excel-google-sheets") + "overlapping-bookings-template.png",
         render(short_sheet_copy(t + "tidy-tabs-overlapping-bookings.xlsx", "Bookings", 14, "bookings"), "Bookings", name="bookings"))
    save(d % (2, "cash-drawer-count-over-short-farmers-market-excel-google-sheets") + "cash-drawer-count-template.png",
         render(t + "tidy-tabs-cash-drawer-count.xlsx", "Count", name="cashdrawer"))
    save(d % (3, "date-of-nth-weekday-in-month-second-saturday-excel-google-sheets") + "nth-weekday-of-month-template.png",
         render(t + "tidy-tabs-nth-weekday-of-month.xlsx", "Dates", name="nthweekday"),
         render(t + "tidy-tabs-nth-weekday-of-month.xlsx", "Schedule", name="nthschedule"))
    save(d % (4, "format-phone-numbers-substitute-left-mid-right-excel-google-sheets") + "format-phone-numbers-template.png",
         render(short_sheet_copy(t + "tidy-tabs-format-phone-numbers.xlsx", "Phones", 14, "phones"), "Phones", name="phones"))
    save(d % (5, "scale-recipe-batch-quantities-absolute-references-excel-google-sheets") + "scale-recipe-batch-template.png",
         render(short_sheet_copy(t + "tidy-tabs-scale-recipe-batch.xlsx", "Recipe", 14, "recipe"), "Recipe", name="recipe"))


def heroes_2026_10_18():
    t = "templates/"
    d = "images/2026-10-18-team-0%d-%s/"
    save(d % (1, "lookup-with-two-criteria-index-match-sumifs-excel-google-sheets") + "lookup-two-criteria-price-lookup-template.png",
         render(short_sheet_copy(t + "tidy-tabs-lookup-two-criteria.xlsx", "Prices", 11, "twocritprices"), "Prices", hide_cols="E", name="twocritprices"),
         render(t + "tidy-tabs-lookup-two-criteria.xlsx", "Lookup", name="twocritlookup"))
    save(d % (2, "stock-on-hand-from-in-out-log-sumifs-excel-google-sheets") + "stock-on-hand-log-template.png",
         render(short_sheet_copy(t + "tidy-tabs-stock-on-hand-log.xlsx", "Stock", 10, "stockonhand"), "Stock", name="stockonhand"))
    save(d % (3, "extract-domain-from-email-address-left-mid-find-excel-google-sheets") + "email-domain-split-template.png",
         render(short_sheet_copy(t + "tidy-tabs-email-domain-split.xlsx", "Emails", 16, "emaildomain"), "Emails", name="emaildomain"))
    save(d % (4, "on-time-shipping-rate-promised-vs-shipped-date-excel-google-sheets") + "on-time-shipping-rate-template.png",
         render(short_sheet_copy(t + "tidy-tabs-on-time-shipping-rate.xlsx", "Orders", 16, "ontimeorders"), "Orders", name="ontimeorders"),
         render(t + "tidy-tabs-on-time-shipping-rate.xlsx", "Summary", name="ontimesummary"))
    save(d % (5, "price-per-square-foot-with-minimum-charge-excel-google-sheets") + "price-per-square-foot-quote-template.png",
         render(short_sheet_copy(t + "tidy-tabs-price-per-square-foot.xlsx", "Quote", 14, "sqftquote"), "Quote", name="sqftquote"))


def heroes_2026_10_19():
    t = "templates/"
    d = "images/2026-10-19-team-0%d-%s/"
    save(d % (1, "highlight-expiring-dates-next-30-days-conditional-formatting-excel-google-sheets") + "expiring-dates-status-template.png",
         render(short_sheet_copy(t + "tidy-tabs-expiring-dates.xlsx", "Items", 21, "expiring"), "Items", name="expiring"))
    save(d % (2, "stop-duplicate-invoice-numbers-data-validation-custom-excel-google-sheets") + "validation-rules-invoice-template.png",
         render(short_sheet_copy(t + "tidy-tabs-validation-rules.xlsx", "Entries", 13, "valentries"), "Entries", name="valentries"),
         render(short_sheet_copy(t + "tidy-tabs-validation-rules.xlsx", "Rules check", 13, "valrules"), "Rules check", hide_cols="FGH", name="valrules"))
    save(d % (3, "sum-with-or-condition-two-channels-sumproduct-excel-google-sheets") + "sum-with-or-two-channels-template.png",
         render(short_sheet_copy(t + "tidy-tabs-sum-with-or.xlsx", "Orders", 25, "oror"), "Orders", name="oror"),
         render(t + "tidy-tabs-sum-with-or.xlsx", "Totals", name="ortotals", widths={"A": 70}))
    save(d % (4, "fill-blank-cells-with-value-above-excel-google-sheets") + "fill-blanks-vendor-report-template.png",
         render(short_sheet_copy(t + "tidy-tabs-fill-blanks-above.xlsx", "Report", 16, "fillrep"), "Report", name="fillrep"),
         render(short_sheet_copy(t + "tidy-tabs-fill-blanks-above.xlsx", "Filled", 16, "fillfilled"), "Filled", hide_cols="BEFGHI", name="fillfilled"))
    save(d % (5, "mailing-labels-3-across-from-address-list-index-row-column-excel-google-sheets") + "mailing-labels-3-across-template.png",
         render(t + "tidy-tabs-mailing-labels.xlsx", "Labels", hide_cols="DEF", name="labels"))


def hide_rows_copy(src, sheet, rows, name):
    """Temp copy of a template tab with the given row numbers hidden."""
    wb = load_workbook(src)
    ws = wb[sheet]
    for r in rows:
        ws.row_dimensions[r].hidden = True
    path = f"{rp.TMP}/{name}-short.xlsx"
    wb.save(path)
    return path


def heroes_2026_10_20():
    t = "templates/"
    d = "images/2026-10-20-team-0%d-%s/"
    save(d % (1, "months-to-pay-off-equipment-balance-nper-excel-google-sheets") + "months-to-pay-off-template.png",
         render(short_sheet_copy(t + "tidy-tabs-months-to-pay-off.xlsx", "Payoff", 10, "payoff"), "Payoff", hide_cols="C", name="payoff"),
         render(short_sheet_copy(t + "tidy-tabs-months-to-pay-off.xlsx", "Schedule", 15, "payoffsched"), "Schedule", name="payoffsched"))
    save(d % (2, "sum-same-cell-across-monthly-sheets-3d-sum-indirect-excel-google-sheets") + "sum-across-monthly-sheets-template.png",
         render(hide_rows_copy(t + "tidy-tabs-sum-across-monthly-sheets.xlsx", "Summary", [5, 6, 7] + list(range(13, 21)) + list(range(22, 25)), "sumsheets"), "Summary", name="sumsheets", widths={"A": 26}))
    save(d % (3, "find-outlier-order-amounts-quartile-iqr-excel-google-sheets") + "outlier-order-amounts-template.png",
         render(hide_rows_copy(t + "tidy-tabs-outlier-order-amounts.xlsx", "Orders", list(range(2, 10)) + list(range(25, 101)), "outliers"), "Orders", name="outliers"),
         render(short_sheet_copy(t + "tidy-tabs-outlier-order-amounts.xlsx", "Fences", 7, "fences"), "Fences", hide_cols="C", name="fences"))
    save(d % (4, "return-on-ad-spend-roas-break-even-excel-google-sheets") + "roas-calculator-template.png",
         render(hide_rows_copy(t + "tidy-tabs-roas-calculator.xlsx", "Campaigns", [13, 14, 15, 16, 18, 19], "roas"), "Campaigns", name="roas"))
    save(d % (5, "pareto-80-20-top-products-cumulative-percent-excel-google-sheets") + "pareto-top-products-template.png",
         render(short_sheet_copy(t + "tidy-tabs-pareto-top-products.xlsx", "Products", 16, "pareto"), "Products", name="pareto"))


def heroes_2026_10_21():
    t = "templates/"
    d = "images/2026-10-21-team-0%d-%s/"
    save(d % (1, "average-annual-sales-growth-cagr-rri-excel-google-sheets") + "annual-growth-rate-template.png",
         render(short_sheet_copy(t + "tidy-tabs-annual-growth-rate.xlsx", "Growth", 10, "growth"), "Growth", hide_cols="F", name="growth"))
    save(d % (2, "count-orders-by-dollar-size-range-countifs-frequency-excel-google-sheets") + "order-size-ranges-template.png",
         render(short_sheet_copy(t + "tidy-tabs-order-size-ranges.xlsx", "Ranges", 8, "sizeranges"), "Ranges", name="sizeranges"))
    save(d % (3, "price-in-effect-on-order-date-price-history-lookup-excel-google-sheets") + "price-in-effect-on-order-date-template.png",
         render(short_sheet_copy(t + "tidy-tabs-price-by-effective-date.xlsx", "Orders", 14, "priceorders"), "Orders", hide_cols="FG", name="priceorders"),
         render(short_sheet_copy(t + "tidy-tabs-price-by-effective-date.xlsx", "Prices", 14, "priceprices"), "Prices", hide_cols="D", name="priceprices"))
    save(d % (4, "how-many-batches-from-materials-on-hand-min-int-excel-google-sheets") + "batches-from-materials-template.png",
         render(short_sheet_copy(t + "tidy-tabs-batches-from-materials.xlsx", "Materials", 15, "batchmat"), "Materials", name="batchmat"))
    save(d % (5, "appointment-schedule-end-times-gaps-buffer-excel-google-sheets") + "appointment-schedule-template.png",
         render(short_sheet_copy(t + "tidy-tabs-appointment-schedule.xlsx", "Schedule", 12, "appts"), "Schedule", hide_cols="HIJKL", name="appts"))


def main():
    if "--only-2026-10-21" in sys.argv:
        heroes_2026_10_21()
        return
    if "--only-2026-10-20" in sys.argv:
        heroes_2026_10_20()
        return
    heroes_2026_10_19()
    if "--new" in sys.argv:
        return
    heroes_2026_10_18()
    if "--new" in sys.argv:
        return
    heroes_2026_10_17()
    heroes_2026_10_16()
    heroes_2026_10_15()
    heroes_2026_10_14()
    heroes_2026_10_13()
    heroes_2026_10_12()
    heroes_2026_10_11()
    return
    # older heroes (rebuilding rewrites their PNGs; the early return above skips them)
    t = "templates/"
    save("images/2026-09-25-a-split-city-state-zip-excel-google-sheets/split-city-state-zip-template.png",
         render(t + "tidy-tabs-split-city-state-zip.xlsx", "Split", name="split"))
    save("images/2026-09-28-a-keep-leading-zeros-zip-codes-excel/zip-codes-leading-zeros-fixed.png",
         render(t + "tidy-tabs-zip-codes-leading-zeros.xlsx", "Broken (numbers)", name="zip"))
    save("images/2026-09-30-a-freelance-project-tracker-google-sheets/freelance-project-tracker-template.png",
         render(t + "tidy-tabs-freelance-project-tracker.xlsx", "Projects", hide_cols="CFGHIKLM", name="proj"))
    save("images/2026-10-01-a-craft-inventory-tracker-google-sheets/craft-inventory-tracker-template.png",
         render(t + "tidy-tabs-craft-inventory-tracker.xlsx", "Items", hide_cols="CDEFGHLNO", name="items"))
    save("images/2026-09-29-team-01-dependent-dropdown-list-excel-google-sheets/dependent-dropdown-lists-tab.png",
         render(t + "tidy-tabs-dependent-dropdown-list.xlsx", "Lists", name="dropdown"))
    save("images/2026-09-29-team-02-client-contact-list-template-google-sheets/client-contact-list-template.png",
         render(t + "tidy-tabs-client-contact-list.xlsx", "Clients", hide_cols="G", name="clients"))
    save("images/2026-09-29-team-03-packing-slip-template-google-sheets/packing-slip-template.png",
         render(t + "tidy-tabs-packing-slip-template.xlsx", "Packing slip", name="packing"))
    save("images/2026-09-29-team-04-running-balance-column-excel-google-sheets/running-balance-cash-log-template.png",
         render(t + "tidy-tabs-running-balance-cash-log.xlsx", "Cash log", name="cashlog"))
    save("images/2026-09-29-team-05-highlight-duplicate-values-excel-google-sheets/highlight-duplicate-emails-template.png",
         render(t + "tidy-tabs-highlight-duplicates-sample.xlsx", "Orders", name="dupes"))
    # 2026-09-30 team posts
    d = "images/2026-09-30-team-0%d-%s/"
    save(d % (1, "receivables-aging-report-excel-google-sheets") + "receivables-aging-report-template.png",
         render(t + "tidy-tabs-receivables-aging-report.xlsx", "Aging", name="aging"))
    save(d % (2, "time-to-decimal-hours-excel-google-sheets") + "time-to-decimal-hours-template.png",
         render(t + "tidy-tabs-time-to-decimal-hours.xlsx", "Shifts", hide_cols="AGI", name="shifts"))
    save(d % (3, "markup-vs-margin-calculator-excel-google-sheets") + "markup-vs-margin-calculator-template.png",
         render(t + "tidy-tabs-markup-vs-margin-calculator.xlsx", "Pricing", hide_cols="EI", name="pricing"))
    save(d % (4, "business-days-ship-by-date-excel-google-sheets") + "business-days-ship-by-date-template.png",
         render(t + "tidy-tabs-business-days-ship-by-date.xlsx", "Orders", hide_cols="BF", name="orders"))
    save(d % (5, "vlookup-price-list-excel-google-sheets") + "vlookup-price-list-template.png",
         render(t + "tidy-tabs-vlookup-price-list.xlsx", "Order lines", hide_cols="FG", name="orderlines"))
    # 2026-10-02 team posts
    d = "images/2026-10-02-team-0%d-%s/"
    save(d % (1, "break-even-calculator-excel-google-sheets") + "break-even-calculator-template.png",
         render(t + "tidy-tabs-break-even-calculator.xlsx", "Break-even", name="breakeven"))
    save(d % (2, "subscription-renewal-tracker-excel-google-sheets") + "subscription-renewal-tracker-template.png",
         render(t + "tidy-tabs-subscription-renewal-tracker.xlsx", "Subscriptions", hide_cols="CDHIJK", name="subs"))
    save(d % (3, "budget-vs-actual-variance-excel-google-sheets") + "budget-vs-actual-variance-template.png",
         render(t + "tidy-tabs-budget-vs-actual.xlsx", "Budget vs actual", hide_cols="GHI", name="budget"))
    save(d % (4, "farmers-market-sales-log-excel-google-sheets") + "farmers-market-sales-log-template.png",
         render(t + "tidy-tabs-farmers-market-sales-log.xlsx", "Sales", name="market"))
    save(d % (5, "clean-customer-list-trim-proper-excel-google-sheets") + "clean-customer-list-trim-proper-template.png",
         render(t + "tidy-tabs-clean-customer-list.xlsx", "Clean names", hide_cols="DE", name="clean"))
    # Print post: the real 6-page default printout next to the real 1-page result.
    d = "images/2026-09-29-a-print-excel-one-page-us-letter/"
    save(d + "print-excel-one-page-before-after.png",
         Image.open(d + "before-six-pages-libreoffice.png").convert("RGB"),
         Image.open(d + "after-one-page-libreoffice.png").convert("RGB"))


if __name__ == "__main__":
    main()
